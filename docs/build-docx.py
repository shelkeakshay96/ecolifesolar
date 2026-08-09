#!/usr/bin/env python3
"""
Render docs/technical-add.md into an MS Word .docx.

    python3 -m venv docs/.venv && docs/.venv/bin/pip install python-docx
    docs/.venv/bin/python docs/build-docx.py

Handles a deliberately small Markdown subset -- exactly what technical-add.md
uses and nothing more:

    #..####          headings 1-4
    paragraphs       with **bold**, *italic*, `inline code`
    - / 1.           bullet and numbered lists
    ```lang ... ```  fenced code -> shaded monospace box
    | a | b |        pipe tables -> real Word tables
    > quote          indented italic with a left rule
    ---              horizontal rule

Fenced blocks are consumed by the line loop before any other rule is tried.
That ordering is load-bearing: the SQL in section 6 contains both `---`-ish
comment lines and `|` characters, and either would otherwise be mistaken for
a horizontal rule or a table row.

The title page and footer are generated from the constants below; everything
else comes from the Markdown, which is the source of truth.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------
# Document identity
# --------------------------------------------------------------------------

VERSION = "2.1"          # bump here only — filename, cover and footer all derive from it

SOURCE = Path(__file__).with_name("technical-add.md")
OUTPUT = Path(__file__).with_name(f"EcoLifeSolar-ADD-Phase1-v{VERSION}.docx")

TITLE = "EcoLifeSolar"
SUBTITLE = "Application Design Document"
TAGLINE = "Phase 1 — Local Development"

COVER_FACTS = [
    ("Version", VERSION),
    ("Status", "Approved for build"),
    ("Date", "9 August 2026"),
    ("Author", "Akshay"),
    ("Client", "Eco Life — Piyush, Satara, Maharashtra"),
    ("Scope", "Phase 1 only. Production deployment deferred; Phase 2 designed, not built."),
    ("Supersedes", "Version 1.0 (flat PHP + Bootstrap), withdrawn in full"),
]

# Palette -- slate neutrals with the solar amber accent, matching section 7.
INK = RGBColor(0x0F, 0x17, 0x2A)
INK_2 = RGBColor(0x1E, 0x29, 0x3B)
INK_3 = RGBColor(0x33, 0x41, 0x55)
INK_4 = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
AMBER = RGBColor(0xD9, 0x77, 0x06)

CODE_BG = "F1F5F9"
CODE_FONT = "Consolas"
BODY_FONT = "Calibri"

HEADING_SPEC = {
    1: (19, INK, True),
    2: (14, INK_2, True),
    3: (12, INK_3, True),
    4: (11, INK_4, True),
}

# --------------------------------------------------------------------------
# Low-level OOXML helpers
# --------------------------------------------------------------------------


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":")), v)
    return e


def shade(element, fill):
    """Apply solid background shading to a cell or paragraph."""
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    pr.append(_el("w:shd", w_val="clear", w_color="auto", w_fill=fill))


def cell_margins(cell, top=80, start=140, bottom=80, end=140):
    mar = _el("w:tcMar")
    for tag, val in (("w:top", top), ("w:start", start), ("w:bottom", bottom), ("w:end", end)):
        mar.append(_el(tag, w_w=str(val), w_type="dxa"))
    cell._tc.get_or_add_tcPr().append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", w_val="true"))


def left_rule(paragraph, color="CBD5E1", size=18):
    bdr = _el("w:pBdr")
    bdr.append(_el("w:left", w_val="single", w_sz=str(size), w_space="8", w_color=color))
    paragraph._p.get_or_add_pPr().append(bdr)


def bottom_rule(paragraph, color="E2E8F0", size=6):
    bdr = _el("w:pBdr")
    bdr.append(_el("w:bottom", w_val="single", w_sz=str(size), w_space="4", w_color=color))
    paragraph._p.get_or_add_pPr().append(bdr)


def field(paragraph, instruction, placeholder=""):
    """Insert a Word field code (TOC, PAGE, NUMPAGES)."""
    run = paragraph.add_run()
    run._r.append(_el("w:fldChar", w_fldCharType="begin"))
    instr = _el("w:instrText", xml_space="preserve")
    instr.text = instruction
    run._r.append(instr)
    run._r.append(_el("w:fldChar", w_fldCharType="separate"))
    text = OxmlElement("w:t")
    text.text = placeholder
    run._r.append(text)
    run._r.append(_el("w:fldChar", w_fldCharType="end"))
    return run


def page_break_before(paragraph):
    paragraph._p.get_or_add_pPr().append(_el("w:pageBreakBefore"))


# --------------------------------------------------------------------------
# Inline formatting
# --------------------------------------------------------------------------

# `code` is matched FIRST so that a span like `**/view/**/*.phtml` is treated as
# code rather than being torn apart by the bold rule.
INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\*[^*\n]+?\*)")

CODE_INK = RGBColor(0xB4, 0x53, 0x09)


def add_inline(paragraph, text, size=None, color=None, italic=False, bold=False):
    """Add text to a paragraph, honouring **bold**, *italic* and `code`.

    Emphasis nests: `**\\`Registry\\`**` renders as bold monospace, because the
    marker handlers recurse rather than emitting their contents verbatim.
    """
    for piece in INLINE.split(text):
        if not piece:
            continue

        if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt((size or 10.5) - 1.0)
            run.font.color.rgb = color or CODE_INK
            run.bold = bold
            run.italic = italic
            continue

        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline(paragraph, piece[2:-2], size, color, italic, bold=True)
            continue

        if piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            add_inline(paragraph, piece[1:-1], size, color, italic=True, bold=bold)
            continue

        run = paragraph.add_run(piece)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return paragraph


# --------------------------------------------------------------------------
# Markdown -> block list
# --------------------------------------------------------------------------


def parse(markdown):
    """Return a list of (kind, payload) blocks."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks, buf, i = [], [], 0

    def flush():
        if buf:
            blocks.append(("p", " ".join(buf).strip()))
            buf.clear()

    while i < len(lines):
        line = lines[i]

        # Fenced code -- consumed first, before any other rule is tried.
        if line.lstrip().startswith("```"):
            flush()
            lang = line.strip().strip("`").strip()
            i += 1
            code = []
            while i < len(lines) and not lines[i].lstrip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            while code and not code[0].strip():
                code.pop(0)
            while code and not code[-1].strip():
                code.pop()
            blocks.append(("code", (lang, code)))
            continue

        stripped = line.strip()

        if not stripped:
            flush()
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush()
            blocks.append(("h", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            flush()
            blocks.append(("hr", None))
            i += 1
            continue

        if stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(("table", rows))
            continue

        if stripped.startswith("> "):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(quote)))
            continue

        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        number = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if bullet or number:
            flush()
            kind = "ul" if bullet else "ol"
            items = []
            while i < len(lines):
                m = re.match(r"^(\s*)[-*]\s+(.*)$", lines[i]) if kind == "ul" \
                    else re.match(r"^(\s*)\d+[.)]\s+(.*)$", lines[i])
                if not m:
                    # continuation line of the previous item
                    if items and lines[i].strip() and lines[i].startswith(("  ", "\t")) \
                            and not lines[i].lstrip().startswith(("-", "*", "|", "#", "```")):
                        level, text = items[-1]
                        items[-1] = (level, text + " " + lines[i].strip())
                        i += 1
                        continue
                    break
                items.append((len(m.group(1)) // 2, m.group(2).strip()))
                i += 1
            blocks.append((kind, items))
            continue

        buf.append(stripped)
        i += 1

    flush()
    return blocks


def split_row(row):
    """Split a pipe-table row, honouring \\| escapes."""
    body = row.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    cells = re.split(r"(?<!\\)\|", body)
    return [c.strip().replace("\\|", "|") for c in cells]


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------


def setup_document():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK_3
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for level, (size, color, bold) in HEADING_SPEC.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(16 if level <= 2 else 12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.2))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"EcoLifeSolar ADD v{VERSION}    ·    Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field(footer, "PAGE", "1")
    run = footer.add_run(" of ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field(footer, "NUMPAGES", "1")
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED

    # Ask Word to refresh the TOC when the document is opened.
    doc.settings.element.append(_el("w:updateFields", w_val="true"))
    return doc


def cover_page(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(TITLE)
    run.font.size = Pt(38)
    run.font.bold = True
    run.font.color.rgb = INK

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(SUBTITLE)
    run.font.size = Pt(19)
    run.font.color.rgb = AMBER

    p = doc.add_paragraph()
    run = p.add_run(TAGLINE)
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    bottom_rule(p)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in COVER_FACTS:
        row = table.add_row()
        for idx, text in enumerate((label, value)):
            cell = row.cells[idx]
            cell.width = Cm(4.2 if idx == 0 else 12.6)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(3)
            run = para.add_run(text)
            run.font.size = Pt(9.5)
            run.bold = idx == 0
            run.font.color.rgb = INK_2 if idx == 0 else INK_3
            cell_margins(cell)
        shade(row.cells[0]._tc, "F8FAFC")

    heading = doc.add_paragraph()
    page_break_before(heading)
    run = heading.add_run("Contents")
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = INK

    note = doc.add_paragraph()
    run = note.add_run(
        "If the list below is empty, select it and press F9 to build it "
        "(Word offers to do this automatically on open)."
    )
    run.font.size = Pt(8.5)
    run.italic = True
    run.font.color.rgb = MUTED

    toc = doc.add_paragraph()
    field(toc, 'TOC \\o "1-3" \\h \\z \\u', "")


def render_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell._tc, CODE_BG)
    cell_margins(cell, top=120, bottom=120, start=160, end=160)

    for index, line in enumerate(lines or [""]):
        para = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        fmt = para.paragraph_format
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)
        fmt.line_spacing = 1.0
        run = para.add_run(line if line.strip() else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = INK_2

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_table(doc, rows):
    parsed = [split_row(r) for r in rows]
    parsed = [r for r in parsed if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
    if not parsed:
        return
    width = max(len(r) for r in parsed)

    table = doc.add_table(rows=0, cols=width)
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        table.style = "Table Grid"

    for row_index, cells in enumerate(parsed):
        row = table.add_row()
        for col in range(width):
            cell = row.cells[col]
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.05
            add_inline(para, cells[col] if col < len(cells) else "", size=9)
            for run in para.runs:
                if row_index == 0:
                    run.bold = True
            cell_margins(cell)
        if row_index == 0:
            repeat_header(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render(doc, blocks):
    stats = {"h": 0, "table": 0, "code": 0, "p": 0, "list": 0}
    first_heading = True

    for kind, payload in blocks:
        if kind == "h":
            level, text = payload
            para = doc.add_paragraph(style=f"Heading {level}")
            add_inline(para, text)
            for run in para.runs:
                run.font.color.rgb = HEADING_SPEC[level][1]
                run.font.size = Pt(HEADING_SPEC[level][0])
                run.font.bold = True
            if level == 1:
                bottom_rule(para)
                if first_heading:
                    first_heading = False
                else:
                    page_break_before(para)
            stats["h"] += 1

        elif kind == "p":
            add_inline(doc.add_paragraph(), payload)
            stats["p"] += 1

        elif kind in ("ul", "ol"):
            style = "List Bullet" if kind == "ul" else "List Number"
            for level, text in payload:
                name = style if level == 0 else f"{style} {min(level + 1, 3)}"
                try:
                    para = doc.add_paragraph(style=name)
                except KeyError:
                    para = doc.add_paragraph(style=style)
                para.paragraph_format.space_after = Pt(3)
                add_inline(para, text)
            stats["list"] += 1

        elif kind == "code":
            render_code(doc, payload[1])
            stats["code"] += 1

        elif kind == "table":
            render_table(doc, payload)
            stats["table"] += 1

        elif kind == "quote":
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(8)
            left_rule(para)
            add_inline(para, payload, color=MUTED, italic=True)

        elif kind == "hr":
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            bottom_rule(para, color="E2E8F0")

    return stats


def main():
    if not SOURCE.exists():
        sys.exit(f"error: {SOURCE} not found")

    blocks = parse(SOURCE.read_text(encoding="utf-8"))
    doc = setup_document()
    cover_page(doc)
    stats = render(doc, blocks)
    doc.save(OUTPUT)

    size_kb = OUTPUT.stat().st_size / 1024
    print(f"wrote {OUTPUT.name}  ({size_kb:.0f} KB)")
    print(
        f"  headings {stats['h']}   paragraphs {stats['p']}   "
        f"tables {stats['table']}   code blocks {stats['code']}   lists {stats['list']}"
    )


if __name__ == "__main__":
    main()
